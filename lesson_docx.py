"""Builds the lesson plan .docx, following the school's two templates:
  - "cbc": Form-2-D-and-T_Lesson-Plans.docx layout (Forms, competences)
  - "old": ALL_BUILDING_LESSONS / HOLDING_TOOLS layout (Grades, LSBAT outcomes)
"""
import os
import re
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

# Change these two if the school changes; or read them from environment variables.
SCHOOL_CBC = os.environ.get("SCHOOL_NAME_CBC", "KENNETH KAUNDA PROVINCIAL STEM SCHOOL")
SCHOOL_OLD = os.environ.get("SCHOOL_NAME_OLD", "KENNETH KAUNDA SECONDARY SCHOOL")
CREST = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "coat_of_arms.png")

# ---- Field definitions (the web form and the document both use these) -------------
CBC_FIELDS = [
    ("general_competences", "General competence(s)"),
    ("lesson_goal", "Lesson goal"),
    ("specific_competences", "Specific competences (one per line)"),
    ("rationale", "Rationale"),
    ("prior_knowledge", "Prior knowledge"),
    ("references", "References"),
    ("learning_environment", "Learning environment"),
    ("materials", "Teaching and learning materials/resources"),
    ("expected_standard", "Expected standard"),
]
CBC_STAGES = [
    ("introduction", "INTRODUCTION"),
    ("development", "LESSON DEVELOPMENT"),
    ("exercise", "EXERCISE/ASSESSMENT"),
    ("homework", "HOMEWORK"),
    ("conclusion", "CONCLUSION"),
]
CBC_STAGE_KEYS = [("teacher", "Teacher's role"), ("learners", "Learners' role"), ("assessment", "Assessment criteria")]

OLD_FIELDS = [
    ("outcomes", "Learning outcomes - LSBAT (one per line)"),
    ("knowledge", "Rationale: Knowledge"),
    ("skills", "Skills"),
    ("values", "Values"),
    ("teaching_aids", "Teaching aids"),
]
OLD_STAGES = [
    ("introduction", "Introduction"),
    ("development", "Development"),
    ("application", "Application"),
    ("conclusion", "Conclusion"),
]
OLD_STAGE_KEYS = [("content", "Lesson content"), ("methodology", "Methodology"), ("activity", "Learner activity")]

META_KEYS = ["teacher", "date", "class_name", "duration", "ts_number", "subject", "topic", "subtopic",
             "boys", "girls", "total"]


def schema():
    return {
        "cbc": {"fields": CBC_FIELDS, "stages": CBC_STAGES, "stage_keys": CBC_STAGE_KEYS},
        "old": {"fields": OLD_FIELDS, "stages": OLD_STAGES, "stage_keys": OLD_STAGE_KEYS},
    }


def stage_times(duration):
    """Split the lesson like the template's 80 min (10 / 50 / 15 / 5)."""
    d = int(duration)
    intro = int(d * 10 / 80 + 0.5)
    exercise = int(d * 15 / 80 + 0.5)
    conclusion = int(d * 5 / 80 + 0.5)
    return {"introduction": intro, "exercise": exercise, "conclusion": conclusion,
            "development": d - intro - exercise - conclusion}


def safe_filename(plan):
    base = f"{plan['meta'].get('subject', 'Lesson')}_{plan['meta'].get('topic', 'Plan')}"
    base = re.sub(r"[^A-Za-z0-9]+", "_", base).strip("_")[:60] or "Lesson_Plan"
    return f"Lesson_Plan_{base}.docx"


# ---- Small docx helpers ------------------------------------------------------------
def _font(run, size=None, bold=None, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _para(container, after=0, align=None):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    return p


def _rich(p, text, size=12, bold=False, name="Times New Roman"):
    """Adds text to paragraph p. Parts wrapped in **double stars** become bold."""
    for i, part in enumerate(re.split(r"\*\*", text or "")):
        if part:
            r = p.add_run(part)
            _font(r, size, bold or (i % 2 == 1), name)


def _label_line(container, label, value, after=0, sep=": "):
    p = _para(container, after)
    r = p.add_run(label)
    _font(r, 12, True)
    r2 = p.add_run(sep + (value or ""))
    _font(r2, 12, False)
    return p


def _two_col_table(container, rows, widths=(2.0, 6.0), style="Table Grid"):
    table = container.add_table(rows=len(rows), cols=2, style=style)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (label, value) in enumerate(rows):
        row = table.rows[i].cells
        _cell_text(row[0], label, bold=True)
        _cell_text(row[1], value)
    _set_widths(table, widths)
    return table


def _cell_text(cell, text, size=12, bold=False):
    """Puts text in a cell: one paragraph per line, **bold** markers supported."""
    lines = [l for l in (text or "").split("\n")] or [""]
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)
        _rich(p, line, size, bold)


def _set_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for col, w in zip(table.columns, widths):
        col.width = Inches(w)
    for row in table.rows:
        _no_split(row)
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _bullets(container, text):
    for line in [l.strip() for l in (text or "").split("\n") if l.strip()]:
        line = re.sub(r"^[-\u2022*]\s*", "", line)
        p = container.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(0)
        _rich(p, line)


def _evaluation(doc):
    _para(doc, 0)
    p = _para(doc, 0)
    _rich(p, "LESSON EVALUATION:", bold=True)
    _label_line(doc, "Challenges", "." * 130, sep=": ").runs[0].bold = False
    _label_line(doc, "Teacher's response", "." * 120, sep=": ").runs[0].bold = False


def _base(page_w, page_h, margins):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(page_w), Inches(page_h)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = [Inches(m) for m in margins]
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return doc


# ---- CBC layout (Forms) ------------------------------------------------------------
def build_cbc(plan):
    m, c = plan["meta"], plan["content"]
    doc = _base(8.5, 11, (0.7, 0.5, 0.4, 0.5))

    if os.path.exists(CREST):
        p = _para(doc, 0, WD_ALIGN_PARAGRAPH.CENTER)
        p.add_run().add_picture(CREST, width=Inches(0.8))
    p = _para(doc, 3, WD_ALIGN_PARAGRAPH.CENTER)
    _rich(p, "MINISTRY OF EDUCATION", bold=True, size=14)
    p = _para(doc, 2, WD_ALIGN_PARAGRAPH.CENTER)
    _rich(p, SCHOOL_CBC.upper(), bold=True, size=16)
    p = _para(doc, 6, WD_ALIGN_PARAGRAPH.CENTER)
    _rich(p, f"{(m['subject'] or 'SUBJECT').upper()} LESSON PLAN", bold=True, size=14)

    p = _para(doc, 0)
    _rich(p, "NAME OF TEACHER: ", bold=True)
    _rich(p, (m["teacher"] or "") + "     ")
    _rich(p, "DATE: ", bold=True)
    _rich(p, m["date"] or "")

    p = _para(doc, 0)
    _rich(p, "CLASS: ", bold=True)
    _rich(p, (m["class_name"] or "") + "     ")
    _rich(p, "DURATION: ", bold=True)
    _rich(p, f"{m['duration']} Minutes     ")
    _rich(p, "TS Number: ", bold=True)
    _rich(p, m["ts_number"] or "")

    _label_line(doc, "SUBJECT:", m["subject"], sep=" ")
    _label_line(doc, "TOPIC:", m["topic"], sep=" ")
    _label_line(doc, "SUB-TOPIC:", m["subtopic"], sep=" ")
    _label_line(doc, "GENERAL COMPETENCE(S):", c["general_competences"], sep=" ")
    _label_line(doc, "LESSON GOAL:", c["lesson_goal"], sep=" ")
    _label_line(doc, "SPECIFIC COMPETENCES:", "", sep="")
    _bullets(doc, c["specific_competences"])
    _label_line(doc, "RATIONALE:", c["rationale"], sep=" ")
    _label_line(doc, "PRIOR KNOWLEDGE:", c["prior_knowledge"], sep=" ")
    _label_line(doc, "REFERENCES:", c["references"], sep=" ")
    _label_line(doc, "LEARNING ENVIRONMENT:", c["learning_environment"], sep=" ")
    _label_line(doc, "TEACHING AND LEARNING MATERIALS/RESOURCES:", c["materials"], sep=" ")
    _label_line(doc, "EXPECTED STANDARD:", c["expected_standard"], sep=" ")

    _para(doc, 0)
    p = _para(doc, 6)
    _rich(p, "LESSON PROGRESSION:", bold=True)

    times = stage_times(m["duration"])
    widths = [2.1, 2.1, 1.5, 1.6]
    t = doc.add_table(rows=1, cols=4, style="Table Grid")
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for cell, head in zip(t.rows[0].cells, ["STAGES", "TEACHER'S ROLE", "LEARNERS' ROLE", "ASSESSMENT CRITERIA"]):
        _cell_text(cell, head, bold=True)
    _repeat_header(t.rows[0])
    for key, label in CBC_STAGES:
        row = t.add_row()
        st = c["stages"][key]
        p = row.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        _rich(p, label, bold=True)
        if key in times:
            _rich(p, f" ({times[key]} mins)")
        _cell_text(row.cells[1], st["teacher"])
        _cell_text(row.cells[2], st["learners"])
        _cell_text(row.cells[3], st["assessment"])
    _set_widths(t, widths)

    _evaluation(doc)
    return doc


# ---- Old curriculum layout (Grades) -------------------------------------------------
def build_old(plan):
    m, c = plan["meta"], plan["content"]
    doc = _base(8.27, 11.69, (0.8, 0.7, 0.7, 0.7))

    p = _para(doc, 4, WD_ALIGN_PARAGRAPH.CENTER)
    _rich(p, f"{SCHOOL_OLD} {m['subject'].upper()} DEPARTMENT", bold=True)
    p = _para(doc, 8, WD_ALIGN_PARAGRAPH.CENTER)
    _rich(p, "LESSON PLAN", bold=True)

    _two_col_table(doc, [
        ("TEACHER", m["teacher"] or ""),
        ("SUBJECT", m["subject"] or ""),
        ("CLASS", m["class_name"] or ""),
        ("DATE", m["date"] or ""),
        ("DURATION", f"{m['duration']} minutes" if m.get("duration") else ""),
        ("BOYS", m["boys"] or ""),
        ("GIRLS", m["girls"] or ""),
        ("TOTAL OF PUPILS", m["total"] or ""),
        ("TOPIC", m["topic"] or ""),
        ("SUBTOPIC", m["subtopic"] or ""),
    ], widths=(2.2, 5.8))

    p = _para(doc, 2)
    _rich(p, "LEARNING OUTCOMES LSBAT", bold=True)
    _bullets(doc, c["outcomes"])

    p = _para(doc, 2)
    p.paragraph_format.space_before = Pt(6)
    _rich(p, "RATIONALE", bold=True)
    p = _para(doc, 8)
    _rich(p, "Knowledge: ", bold=True)
    _rich(p, c["knowledge"])

    widths = [1.2, 2.6, 1.7, 1.25]
    t = doc.add_table(rows=1, cols=4, style="Table Grid")
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for cell, head in zip(t.rows[0].cells, ["PART", "LESSON CONTENT", "METHODOLOGY", "LEARNER ACTIVITY"]):
        _cell_text(cell, head, bold=True)
    _repeat_header(t.rows[0])
    for key, label in OLD_STAGES:
        row = t.add_row()
        st = c["stages"][key]
        _cell_text(row.cells[0], label)
        _cell_text(row.cells[1], st["content"])
        _cell_text(row.cells[2], st["methodology"])
        _cell_text(row.cells[3], st["activity"])
    _set_widths(t, widths)

    p = _para(doc, 2)
    p.paragraph_format.space_before = Pt(8)
    _rich(p, "Skills: ", bold=True)
    _rich(p, c["skills"])
    p = _para(doc, 2)
    _rich(p, "Values: ", bold=True)
    _rich(p, c["values"])
    p = _para(doc, 2)
    _rich(p, "TEACHING AIDS: ", bold=True)
    _rich(p, c["teaching_aids"])

    _evaluation(doc)
    return doc


def build_docx(plan):
    doc = build_cbc(plan) if plan["layout"] == "cbc" else build_old(plan)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
